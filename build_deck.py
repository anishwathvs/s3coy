"""Builds the D3COY 1.0 Round 0 deck (team CTRL FREAKS) as a fully editable PPTX.

The official template artwork (titles, logos, frames) is kept untouched as each
slide's background; every piece of content on top is a native, editable
PowerPoint shape, text box, picture or chart.

Usage: python3 build_deck.py <assets_dir> <out.pptx>
  assets_dir/bg/slide1.jpg .. slide6.jpg   template backgrounds (4000x2250)
  assets_dir/icons/*.png                    Lucide icons + logo_*.png tech logos
"""
import copy
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ASSETS = Path(sys.argv[1])
OUT = Path(sys.argv[2])
ICONS = ASSETS / "icons"

HEAD = "Georgia"
BODY = "Calibri"

RED = "C62828"
RED_HI = "EF5350"
RED_DK = "5C1015"
CARD = "160506"
CARD2 = "240A0D"
LINE = "8E2A30"
CREAM = "F7ECDD"
MUTED = "D6C3B0"
GOLD = "E8B04B"
GREEN = "43A047"
AMBER = "F9A825"


# ---------------------------------------------------------------- helpers
def rgb(h):
    return RGBColor.from_string(h)


def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _alpha(fill_parent, pct):
    clr = fill_parent.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(pct * 1000)))


def box(s, x, y, w, h, fill=CARD, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.08, alpha=None, dash=False):
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(shp)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _alpha(shp._element.spPr.find(qn("a:solidFill")), alpha)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dash:
            ln = shp._element.spPr.find(qn("a:ln"))
            d = etree.SubElement(ln, qn("a:prstDash"))
            d.set("val", "dash")
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    shp.text_frame.text = ""
    return shp


def text(s, x, y, w, h, paras, size=11, color=CREAM, bold=False, italic=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, after=0, margin=0.04):
    """paras: str | list of paragraphs; a paragraph is str or list of runs;
    a run is str or (str, {opts}) with opts size/color/bold/italic/font."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        para.space_after = Pt(after)
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = t
            f = run.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def icon(s, name, x, y, size):
    return s.shapes.add_picture(str(ICONS / f"{name}.png"), Inches(x), Inches(y),
                                Inches(size), Inches(size))


def badge(s, name, x, y, d, fill=RED, ring=None, pad=0.22):
    c = box(s, x, y, d, d, fill=fill, line=ring, lw=1.5, shape=MSO_SHAPE.OVAL)
    icon(s, name, x + d * pad, y + d * pad, d * (1 - 2 * pad))
    return c


def num_badge(s, n, x, y, d, fill=RED, size=12):
    c = box(s, x, y, d, d, fill=fill, line=CREAM, lw=1.25, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = HEAD
    r.font.color.rgb = rgb(CREAM)
    return c


def arrow(s, x1, y1, x2, y2, color=GOLD, w=2.0, head=True, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


def section_title(s, x, y, w, label, size=14, color=CREAM, icon_name=None, align=PP_ALIGN.LEFT,
                  underline=True):
    tx = x
    if icon_name:
        badge(s, icon_name, x, y + 0.01, 0.34, pad=0.2)
        tx = x + 0.42
    text(s, tx, y, w - (tx - x), 0.36, [[(label, {"bold": True})]], size=size, color=color,
         font=HEAD, align=align, anchor=MSO_ANCHOR.MIDDLE)
    if underline:
        arrow(s, x, y + 0.4, x + w, y + 0.4, color=RED, w=1.25, head=False)


def pill(s, x, y, w, h, label, fill=RED_DK, line=RED, size=10, color=CREAM, bold=True,
         align=PP_ALIGN.CENTER):
    b = box(s, x, y, w, h, fill=fill, line=line, lw=1.0, radius=h / 2)
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = BODY
    r.font.color.rgb = rgb(color)
    return b


def set_background(slide, img):
    _, rid = slide.part.get_or_add_image_part(str(img))
    cSld = slide._element.find(qn("p:cSld"))
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    bg = etree.SubElement(cSld, qn("p:bg"))
    cSld.remove(bg)
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, qn("p:bgPr"))
    bf = etree.SubElement(bgPr, qn("a:blipFill"))
    bf.set("dpi", "0")
    bf.set("rotWithShape", "1")
    blip = etree.SubElement(bf, qn("a:blip"))
    blip.set(qn("r:embed"), rid)
    etree.SubElement(bf, qn("a:srcRect"))
    st = etree.SubElement(bf, qn("a:stretch"))
    etree.SubElement(st, qn("a:fillRect"))
    etree.SubElement(bgPr, qn("a:effectLst"))


def card(s, x, y, w, h, title=None, icon_name=None, fill=CARD, line=LINE, alpha=88,
         title_size=12.5, head_fill=RED_DK):
    """Dark card with optional title strip. Returns y where body content starts."""
    box(s, x, y, w, h, fill=fill, line=line, lw=1.0, alpha=alpha)
    if not title:
        return y + 0.08
    hh = 0.4
    box(s, x, y, w, hh, fill=head_fill, line=line, lw=1.0, radius=0.08)
    tx = x + 0.1
    if icon_name:
        badge(s, icon_name, x + 0.07, y + 0.05, 0.3, pad=0.2)
        tx = x + 0.43
    text(s, tx, y, w - (tx - x) - 0.06, hh, [[(title, {"bold": True})]], size=title_size,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    return y + hh + 0.06


def bullets(s, x, y, w, h, items, size=10.5, color=CREAM, bullet="•", bcolor=RED_HI, after=3,
            spacing=1.0):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            head, rest = it
            paras.append([(bullet + " ", {"color": bcolor, "bold": True}),
                          (head, {"bold": True}), (rest, {"color": MUTED})])
        else:
            paras.append([(bullet + " ", {"color": bcolor, "bold": True}), it])
    return text(s, x, y, w, h, paras, size=size, color=color, after=after, spacing=spacing)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide(n):
    sl = prs.slides.add_slide(BLANK)
    set_background(sl, ASSETS / "bg" / f"slide{n}.jpg")
    return sl


# =========================================================== 1. TITLE SLIDE
s = new_slide(1)
# Template frames (inches): team 1.07-4.31 x 1.87-3.10, college 4.43-9.54 x 1.87-3.10,
# problem 1.07-9.54 x 3.22-4.42, members 1.07-9.54 x 4.53-7.06. Labels sit in the top ~0.35".
text(s, 1.15, 2.28, 3.08, 0.72, [[("CTRL FREAKS", {"bold": True})]], size=26, font=HEAD,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 4.5, 2.24, 4.97, 0.8,
     [[("CHENNAI INSTITUTE OF TECHNOLOGY", {"bold": True})],
      [("Chennai, Tamil Nadu", {"size": 11, "color": MUTED, "font": BODY})]],
     size=14.5, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=2)
text(s, 1.17, 3.56, 8.28, 0.82,
     [[("Identity Drift Detection in Autonomous AI Agents", {"bold": True, "font": HEAD,
                                                              "size": 13, "color": GOLD})],
      "AI agents carry their own tokens and permissions. Access control checks whether an "
      "identity is allowed to act — never whether it still behaves like its role. A hidden "
      "prompt, a poisoned tool or a stolen token can hijack an agent while every check passes."],
     size=11, after=2, anchor=MSO_ANCHOR.MIDDLE)

members = [("KAMESHWARAN B", "+91 72008 05541"), ("VS ANISHWATH", "+91 88255 14383"),
           ("LIKITH KRISHNA J", "+91 92823 26251"), ("SRI AARUSH TEJ", "+91 86376 61221")]
row_y, row_h = 4.98, 0.5
for i, (name, phone) in enumerate(members):
    y = row_y + i * row_h
    box(s, 1.25, y + 0.04, 8.1, row_h - 0.08, fill=CARD, line=RED_DK, lw=0.75, alpha=70,
        radius=0.06)
    num_badge(s, f"0{i + 1}", 1.35, y + 0.08, 0.34, size=10)
    text(s, 1.85, y, 4.3, row_h, [[(name, {"bold": True})]], size=15, font=HEAD,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.4, y, 2.85, row_h, phone, size=14, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
         color=MUTED)


# ===================================================== 2. PROPOSED SOLUTION
s = new_slide(2)
TOP = 1.72

# ---- left column: problems faced
box(s, 0.32, TOP, 3.88, 7.06 - TOP, fill=CARD, line=LINE, alpha=90)
section_title(s, 0.45, TOP + 0.06, 3.62, "PROBLEMS FACED", size=14, icon_name="triangle-alert")
threats = [
    ("file-warning", "Prompt Injection", "hidden instruction in a document",
     "Agent leaks data — token still valid"),
    ("plug-zap", "Poisoned Tool", "malicious plugin / API response",
     "Tool calls redirected to attacker"),
    ("key-round", "Stolen Token", "credentials replayed by attacker",
     "Attacker acts as a trusted agent"),
]
ty = TOP + 0.6
for ic, t1, t2, res in threats:
    box(s, 0.45, ty, 3.62, 0.98, fill=CARD2, line=RED_DK, lw=0.75, alpha=95)
    badge(s, ic, 0.55, ty + 0.12, 0.5)
    text(s, 1.1, ty + 0.06, 2.9, 0.3, [[(t1, {"bold": True})]], size=12.5, font=HEAD)
    text(s, 1.1, ty + 0.33, 2.9, 0.25, t2, size=10, color=MUTED, italic=True)
    arrow(s, 0.8, ty + 0.66, 1.08, ty + 0.75, color=GOLD, w=1.5)
    pill(s, 1.12, ty + 0.62, 2.85, 0.28, res, fill="7A1A20", line=RED_HI, size=9.5)
    ty += 1.06
section_title(s, 0.45, ty + 0.04, 3.62, "Limitations of existing solutions", size=12,
              color=GOLD, underline=False)
bullets(s, 0.5, ty + 0.42, 3.6, 1.6, [
    ("IAM checks the token, ", "never the behaviour"),
    ("Static RBAC ", "can't see a role drifting"),
    ("Anomaly tools ", "give black-box scores"),
    ("Kill-switch response ", "halts the whole fleet"),
], size=10.5, after=2)

# ---- right: solution pipeline
RX = 4.35
text(s, RX, TOP - 0.02, 8.6, 0.36,
     [[("SentinelTrace ", {"bold": True, "color": GOLD, "font": HEAD}),
       ("watches what an agent ", {}), ("does", {"bold": True, "italic": True}),
       (" — not just what it is allowed to do.", {})]],
     size=13, italic=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)

CW, GAP = 1.94, 0.28
xs = [RX + i * (CW + GAP) for i in range(4)]
r1y, r1h = TOP + 0.42, 1.4
stages = [
    ("bot", "Agent Fleet", ["Agents A · B · C", "Docs, tools & APIs", "Signed role contract", "Owner-approved scope"]),
    ("scan-eye", "Capture", ["Every action logged", "Trusted / untrusted tag", "Provenance lineage", "No agent code change"]),
    ("brain-circuit", "Behaviour AI", ["Per-role sequence model", "Embedding similarity", "Resource-access graph", "Rolling-window baseline"]),
    ("gauge", "Drift Score", ["Role conformance", "Behavioural novelty", "Untrusted-content signal", "Fused into 0 – 100"]),
]
fills = ["4A1418", "5C1015", "6E1218", "8B1A1F"]
for i, (ic, t, items) in enumerate(stages):
    by = card(s, xs[i], r1y, CW, r1h, title=t, icon_name=ic, head_fill=fills[i], title_size=12)
    bullets(s, xs[i] + 0.06, by, CW - 0.1, r1h - 0.5, items, size=10, after=1)
    if i < 3:
        arrow(s, xs[i] + CW + 0.02, r1y + r1h / 2, xs[i + 1] - 0.02, r1y + r1h / 2, w=2.25)

r2y, r2h = r1y + r1h + 0.3, 1.2
arrow(s, xs[3] + CW / 2, r1y + r1h + 0.02, xs[3] + CW / 2, r2y - 0.02, w=2.25)
row2 = [
    ("layout-dashboard", "Dashboard", "Live scores, reasons & full audit trail"),
    ("shield-check", "Containment", None),
    ("bell-ring", "Explain", "Names the exact action, mismatch and trigger event"),
    ("network", "Chain Risk", "Agent A calls B → B inherits a share of A's risk"),
]
for i, (ic, t, desc) in enumerate(row2):
    by = card(s, xs[i], r2y, CW, r2h, title=t, icon_name=ic, head_fill=fills[3 - i], title_size=12)
    if desc:
        text(s, xs[i] + 0.08, by, CW - 0.14, r2h - 0.48, desc, size=10.5, color=CREAM)
    else:
        for j, (lab, col) in enumerate([("LOG", GREEN), ("APPROVE", AMBER), ("ISOLATE", RED_HI)]):
            pill(s, xs[i] + 0.12, by + 0.02 + j * 0.27, CW - 0.24, 0.23, lab, fill=CARD,
                 line=col, color=col, size=9)
    if i < 3:
        arrow(s, xs[i + 1] - 0.02, r2y + r2h / 2, xs[i] + CW + 0.02, r2y + r2h / 2, w=2.25)

# ---- bottom row: live chart, worked example, novelty
by0 = r2y + r2h + 0.14
bh = 7.06 - by0
# live drift chart (native, editable chart)
box(s, RX, by0, 2.72, bh, fill=CARD, line=LINE, alpha=90)
text(s, RX + 0.1, by0 + 0.02, 2.6, 0.28, [[("Live Drift Score — Agent B", {"bold": True})]],
     size=10.5, font=HEAD, color=GOLD)
cd = CategoryChartData()
cd.categories = [f"t{i}" for i in range(1, 13)]
cd.add_series("Agent B", (12, 14, 11, 15, 13, 16, 22, 48, 71, 88, 34, 18))
cd.add_series("Isolate ≥ 70", (70,) * 12)
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(RX + 0.02), Inches(by0 + 0.26),
                        Inches(2.68), Inches(bh - 0.3), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(8)
ch.legend.font.color.rgb = rgb(MUTED)
ch.font.size = Pt(8)
ch.font.color.rgb = rgb(MUTED)
va = ch.value_axis
va.maximum_scale, va.minimum_scale, va.major_unit = 100, 0, 25
va.major_gridlines.format.line.color.rgb = rgb("3A1A1D")
va.format.line.fill.background()
ca = ch.category_axis
ca.format.line.color.rgb = rgb(LINE)
ca.tick_labels.font.size = Pt(7)
p0, p1 = ch.plots[0].series
p0.format.line.color.rgb = rgb(RED_HI)
p0.format.line.width = Pt(2.5)
p0.smooth = True
p1.format.line.color.rgb = rgb(GOLD)
p1.format.line.width = Pt(1.25)
p1.format.line.dash_style = 4  # dash
for el in (gf.chart._chartSpace,):
    sp = el.find(qn("c:spPr"))
    if sp is None:
        sp = etree.SubElement(el, qn("c:spPr"))
    etree.SubElement(sp, qn("a:noFill"))
    ln = etree.SubElement(sp, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))

# worked example
ex = RX + 2.84
box(s, ex, by0, 2.72, bh, fill=CARD, line=LINE, alpha=90)
text(s, ex + 0.1, by0 + 0.02, 2.6, 0.28, [[("Attack caught in 3 steps", {"bold": True})]],
     size=10.5, font=HEAD, color=GOLD)
steps = [("Reads untrusted PDF", "provenance: untrusted"),
         ("Calls export_data()", "outside role contract"),
         ("Sends to unknown domain", "new access-graph edge")]
sy = by0 + 0.36
for i, (a, b) in enumerate(steps):
    num_badge(s, i + 1, ex + 0.12, sy + 0.02, 0.28, size=10)
    text(s, ex + 0.48, sy - 0.03, 2.2, 0.24, [[(a, {"bold": True})]], size=10)
    text(s, ex + 0.48, sy + 0.17, 2.2, 0.2, b, size=8.5, color=MUTED, italic=True)
    sy += 0.4
pill(s, ex + 0.12, by0 + bh - 0.34, 2.48, 0.26, "Score 88  →  auto-isolated", fill="7A1A20",
     line=RED_HI, size=10)

# novelty
nx = ex + 2.84
nw = 12.98 - nx
box(s, nx, by0, nw, bh, fill=CARD, line=LINE, alpha=90)
text(s, nx + 0.1, by0 + 0.02, nw - 0.2, 0.28, [[("Novelty & Uniqueness", {"bold": True})]],
     size=10.5, font=HEAD, color=GOLD)
nov = ["Behaviour, not just tokens", "Chain-aware risk propagation",
       "Plain-language explanations", "Graduated, not kill-switch"]
pw = (nw - 0.3) / 2
for i, t in enumerate(nov):
    px = nx + 0.1 + (i % 2) * (pw + 0.1)
    py = by0 + 0.36 + (i // 2) * ((bh - 0.44) / 2)
    b = box(s, px, py, pw, (bh - 0.52) / 2, fill="3D0F14", line=RED, lw=1.0, radius=0.1)
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    r.font.size, r.font.bold, r.font.name = Pt(10), True, BODY
    r.font.color.rgb = rgb(CREAM)


# ===================================================== 3. TECHNICAL APPROACH
s = new_slide(3)
ZX, ZW, ZG = 0.32, 2.3, 0.14
zx = [ZX + i * (ZW + ZG) for i in range(4)]
zy, zh = 1.72, 4.3
zones = [
    ("ZONE 1", "Agent Ecosystem", "4A1418"),
    ("ZONE 2", "Telemetry Layer", "5C1015"),
    ("ZONE 3", "Detection Engine", "6E1218"),
    ("ZONE 4", "Scoring & Response", "8B1A1F"),
]
for i, (z, t, f) in enumerate(zones):
    box(s, zx[i], zy, ZW, zh, fill=CARD, line=LINE, alpha=88, dash=True)
    box(s, zx[i], zy, ZW, 0.56, fill=f, line=LINE, radius=0.08)
    text(s, zx[i], zy + 0.02, ZW, 0.54,
         [[(z, {"size": 9, "color": GOLD, "bold": True})], [(t, {"bold": True})]],
         size=12, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def zbox(i, y, h, ic, title, sub, fill=CARD2, line=RED_DK):
    x = zx[i] + 0.1
    w = ZW - 0.2
    box(s, x, y, w, h, fill=fill, line=line, lw=1.0, radius=0.07)
    badge(s, ic, x + 0.07, y + (h - 0.36) / 2, 0.36, pad=0.2)
    text(s, x + 0.48, y, w - 0.52, h,
         [[(title, {"bold": True, "size": 10.5})], [(sub, {"color": MUTED, "size": 9})]],
         anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)


# zone 1
cy = zy + 0.66
for name, role in [("Agent A", "Research"), ("Agent B", "Finance"), ("Agent C", "Ops")]:
    zbox(0, cy, 0.5, "bot", name, f"role: {role}")
    cy += 0.56
box(s, zx[0] + 0.1, cy + 0.02, ZW - 0.2, 0.92, fill="1E0B0D", line=GOLD, lw=0.75, radius=0.07,
    dash=True)
text(s, zx[0] + 0.14, cy + 0.04, ZW - 0.28, 0.25, [[("Inputs", {"bold": True})]], size=9.5,
     color=GOLD, align=PP_ALIGN.CENTER)
for j, (ic, lab) in enumerate([("file-text", "Docs"), ("wrench", "Tools"), ("globe", "APIs")]):
    ix = zx[0] + 0.22 + j * 0.68
    icon(s, ic, ix + 0.1, cy + 0.32, 0.34)
    text(s, ix - 0.05, cy + 0.66, 0.64, 0.22, lab, size=9, align=PP_ALIGN.CENTER, color=MUTED)
cy += 1.02
zbox(0, cy, 0.62, "scroll-text", "Role Contract", "signed: tools, data, flows", fill="3D0F14",
     line=GOLD)

# zone 2
z2 = [("webhook", "Action Interceptor", "SDK hook / proxy on every call"),
      ("database", "Event Log", "time-stamped JSONL stream"),
      ("tags", "Provenance Tagger", "trusted vs untrusted source"),
      ("filter", "Feature Builder", "n-gram sequences + embeddings")]
cy = zy + 0.66
for ic, t, sub in z2:
    zbox(1, cy, 0.8, ic, t, sub)
    cy += 0.88

# zone 3
z3 = [("scroll-text", "Role Conformance", "action vs signed contract"),
      ("activity", "Sequence Anomaly", "Isolation Forest / OC-SVM"),
      ("brain-circuit", "Semantic Drift", "sentence-transformers + FAISS"),
      ("share-2", "Access Graph", "new edge = novelty"),
      ("list-checks", "Rule Layer", "sensitive act ≤ N steps after untrusted read")]
cy = zy + 0.66
for ic, t, sub in z3:
    zbox(2, cy, 0.65, ic, t, sub)
    cy += 0.71

# zone 4
cy = zy + 0.66
box(s, zx[3] + 0.1, cy, ZW - 0.2, 1.0, fill="3D0F14", line=GOLD, lw=1.0, radius=0.07)
text(s, zx[3] + 0.14, cy + 0.02, ZW - 0.28, 0.98,
     [[("Risk Fusion (0–100)", {"bold": True, "size": 10.5, "color": GOLD})],
      [("R = 0.40·Role + 0.35·Novelty + 0.25·Provenance", {"size": 9, "font": HEAD})],
      [("+ α · upstream risk (propagation)", {"size": 9, "color": MUTED})]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, after=1)
cy += 1.08
box(s, zx[3] + 0.1, cy, ZW - 0.2, 1.22, fill=CARD2, line=RED_DK, radius=0.07)
text(s, zx[3] + 0.14, cy + 0.02, ZW - 0.28, 0.26,
     [[("OPA Policy Bands", {"bold": True})]], size=10.5, align=PP_ALIGN.CENTER)
for j, (lab, col) in enumerate([("0–39   LOG", GREEN), ("40–69   APPROVE", AMBER),
                                ("70–100   ISOLATE", RED_HI)]):
    pill(s, zx[3] + 0.22, cy + 0.32 + j * 0.29, ZW - 0.44, 0.24, lab, fill=CARD, line=col,
         color=col, size=9)
cy += 1.3
zbox(3, cy, 0.62, "layout-dashboard", "Streamlit Dashboard", "live scores + reasons")
cy += 0.68
zbox(3, cy, 0.56, "file-check", "Audit Log", "every score & signal kept")

# arrows between zones with labels
for i, lab in enumerate(["actions", "features", "signals"]):
    x1 = zx[i] + ZW + 0.005
    x2 = zx[i + 1] - 0.005
    yy = zy + 2.2
    arrow(s, x1, yy, x2, yy, w=2.25)
# feedback (response back to agents)
arrow(s, zx[3] + ZW / 2, zy + zh + 0.02, zx[3] + ZW / 2, zy + zh + 0.16, color=RED_HI, w=1.5,
      head=False)
arrow(s, zx[3] + ZW / 2, zy + zh + 0.16, zx[0] + ZW / 2, zy + zh + 0.16, color=RED_HI, w=1.5,
      head=False, dash=True)
arrow(s, zx[0] + ZW / 2, zy + zh + 0.16, zx[0] + ZW / 2, zy + zh + 0.02, color=RED_HI, w=1.5)
pill(s, zx[1] + 0.3, zy + zh + 0.05, 4.4, 0.22,
     "containment fed back to the agent: log · pause for approval · sandbox", fill="1E0B0D",
     line=RED_HI, color=RED_HI, size=8.5, bold=False)

# implementation process (right column)
IX = zx[3] + ZW + 0.16
IW = 12.98 - IX
box(s, IX, zy, IW, zh + 0.3, fill=CARD, line=LINE, alpha=90)
text(s, IX + 0.1, zy + 0.04, IW - 0.2, 0.36, [[("Implementation Process", {"bold": True})]],
     size=13, font=HEAD, color=GOLD, align=PP_ALIGN.CENTER)
proc = [("flask-conical", "Simulate", "3–5 agents · normal + injected-attack logs"),
        ("brain-circuit", "Baseline", "fit per-role sequence & access models"),
        ("sigma", "Score", "fuse 3 signals + rule layer → 0–100"),
        ("shield-check", "Respond", "OPA bands + risk propagation"),
        ("rocket", "Demo", "clean run vs attack run, side by side")]
py0 = zy + 0.5
step = (zh + 0.3 - 0.6) / 5
arrow(s, IX + 0.37, py0 + 0.25, IX + 0.37, py0 + 4 * step + 0.25, color=RED, w=2, head=False)
for i, (ic, t, sub) in enumerate(proc):
    yy = py0 + i * step
    badge(s, ic, IX + 0.12, yy, 0.5, ring=CREAM)
    text(s, IX + 0.7, yy - 0.04, IW - 0.78, 0.28,
         [[(f"{i + 1}. ", {"color": GOLD}), (t, {})]], size=11.5, bold=True, font=HEAD)
    text(s, IX + 0.7, yy + 0.22, IW - 0.78, 0.5, sub, size=9.5, color=MUTED)

# tech stack band
ty0 = 6.3
box(s, 0.32, ty0, 12.66, 0.76, fill=CARD, line=LINE, alpha=92)
box(s, 0.32, ty0, 1.45, 0.76, fill=RED_DK, line=LINE, radius=0.08)
text(s, 0.36, ty0, 1.4, 0.76, [[("TECH", {"bold": True})], [("STACK", {"bold": True})]],
     size=12, font=HEAD, color=GOLD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
     spacing=0.9)
stack = [("logo_python", "Python", "core"), ("logo_scikitlearn", "sklearn", "IF / OC-SVM"),
         ("activity", "PyOD", "outliers"), ("logo_huggingface", "SBERT", "embeddings"),
         ("logo_meta", "FAISS", "vector search"), ("scale", "OPA", "policy"),
         ("logo_plotly", "Plotly", "charts"),
         ("logo_streamlit", "Streamlit", "dashboard")]
sw = (12.98 - 1.9) / len(stack)
for i, (ic, a, b) in enumerate(stack):
    x = 1.88 + i * sw
    box(s, x, ty0 + 0.13, 0.5, 0.5, fill=CREAM if ic.startswith("logo_") else RED,
        line=None, radius=0.1)
    icon(s, ic, x + 0.06, ty0 + 0.19, 0.38)
    text(s, x + 0.54, ty0 + 0.1, sw - 0.56, 0.56,
         [[(a, {"bold": True})], [(b, {"color": MUTED, "size": 8.5})]], size=10,
         anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)


# ================================================ 4. FEASIBILITY & VIABILITY
s = new_slide(4)
TY, TH = 1.72, 2.86
# left: feasibility analysis
box(s, 0.32, TY, 6.4, TH, fill=CARD, line=LINE, alpha=90)
cx, cyc, d = 0.5, TY + (TH - 1.95) / 2, 1.95
box(s, cx, cyc, d, d, fill="2A0B0F", line=RED, lw=5, shape=MSO_SHAPE.OVAL)
text(s, cx, cyc, d, d, [[("Feasibility", {"bold": True})], [("Analysis", {"bold": True})]],
     size=16, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
feas = [("cpu", "Technical Feasibility",
         "Proven open-source ML only — Isolation Forest, One-Class SVM, SBERT, FAISS. "
         "No custom training, no GPU."),
        ("indian-rupee", "Economical Feasibility",
         "₹0 licence cost — 100% open-source. Prototype runs on one CPU-only laptop "
         "(8 GB RAM); a pilot needs one small VM."),
        ("clipboard-check", "Operational Feasibility",
         "Sidecar beside existing IAM — agents only need a logging hook. Built in 24 hrs "
         "with 3–5 simulated agents.")]
fy = TY + 0.1
fh = (TH - 0.2 - 0.16) / 3
for ic, t, body in feas:
    fx = 2.85
    arrow(s, cx + d - 0.08, cyc + d / 2, fx - 0.02, fy + fh / 2, color=GOLD, w=1.5, head=False)
    box(s, fx, fy, 3.78, fh, fill=CARD2, line=RED_DK, radius=0.08)
    badge(s, ic, fx + 0.08, fy + 0.08, 0.38, pad=0.2)
    text(s, fx + 0.52, fy + 0.04, 3.2, 0.3, [[(t, {"bold": True})]], size=11.5, font=HEAD,
         color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, fx + 0.52, fy + 0.32, 3.2, fh - 0.34, body, size=9.5)
    fy += fh + 0.08

# right: viability / scalability
VX = 6.86
box(s, VX, TY, 12.98 - VX, TH, fill=CARD, line=LINE, alpha=90)
vcx = 12.98 - 0.18 - d
box(s, vcx, cyc, d, d, fill=RED, line="FFCDD2", lw=3, shape=MSO_SHAPE.OVAL)
text(s, vcx, cyc, d, d, [[("Viability /", {"bold": True})], [("Scalability", {"bold": True})]],
     size=16, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
via = [("server", "Deployment Viability",
        "Docker sidecar that scales out per fleet; one hook plugs into any agent framework."),
       ("coins", "Economic Viability",
        "No cloud lock-in, no per-seat licences; CPU-only scoring stays cheap at scale."),
       ("lock", "Security Viability",
        "Logs stay on-prem, only drift patterns leave; every decision is audit-logged.")]
fy = TY + 0.1
for ic, t, body in via:
    fx = VX + 0.12
    fw = vcx - 0.25 - fx
    arrow(s, fx + fw + 0.02, fy + fh / 2, vcx + 0.08, cyc + d / 2, color=GOLD, w=1.5, head=False)
    box(s, fx, fy, fw, fh, fill=CARD2, line=RED_DK, radius=0.08)
    badge(s, ic, fx + 0.08, fy + 0.08, 0.38, pad=0.2)
    text(s, fx + 0.52, fy + 0.04, fw - 0.6, 0.3, [[(t, {"bold": True})]], size=11.5, font=HEAD,
         color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, fx + 0.52, fy + 0.32, fw - 0.58, fh - 0.34, body, size=9.5)
    fy += fh + 0.08

# bottom-left: challenges & mitigations
BY = TY + TH + 0.14
BH = 7.06 - BY
box(s, 0.32, BY, 6.4, BH, fill=CARD, line=LINE, alpha=90)
text(s, 0.45, BY + 0.04, 2.9, 0.32, [[("Challenges", {"bold": True})]], size=13, font=HEAD,
     color=RED_HI, align=PP_ALIGN.CENTER)
text(s, 3.65, BY + 0.04, 2.95, 0.32, [[("Mitigations", {"bold": True})]], size=13, font=HEAD,
     color=GREEN, align=PP_ALIGN.CENTER)
cm = [("Cold start for brand-new agents", "Role-template baseline pooled from similar agents"),
      ("False positives stop real work", "Graduated response: log → approve → isolate"),
      ("Legitimate behaviour evolves", "Rolling-window re-baselining keeps models current"),
      ("One weak signal is noisy", "Multi-signal fusion — no single signal decides")]
rh = (BH - 0.46) / 4
for i, (c, m) in enumerate(cm):
    yy = BY + 0.4 + i * rh
    pill(s, 0.45, yy, 2.9, rh - 0.08, c, fill="4A1418", line=RED_HI, size=9.5)
    arrow(s, 3.38, yy + (rh - 0.08) / 2, 3.62, yy + (rh - 0.08) / 2, color=GOLD, w=2)
    pill(s, 3.65, yy, 2.95, rh - 0.08, m, fill="1B3A1E", line=GREEN, size=9.5)

# bottom-right: SWOT
box(s, VX, BY, 12.98 - VX, BH, fill=CARD, line=LINE, alpha=90)
text(s, VX, BY + 0.02, 12.98 - VX, 0.32, [[("SWOT ANALYSIS", {"bold": True})]], size=13,
     font=HEAD, color=GOLD, align=PP_ALIGN.CENTER)
swot = [("S", "Strengths", "badge-check", "2E7D32",
         ["Identity + behaviour in one score", "Explainable, fully open-source", "Works with any agent framework"]),
        ("W", "Weaknesses", "hourglass", "F57F17",
         ["Needs history to baseline", "Demo uses synthetic logs", "Weights need per-org tuning"]),
        ("O", "Opportunities", "trending-up", "1565C0",
         ["Prompt injection is #1 in OWASP Top 10 for LLM apps", "Agent fleets growing fast", "New AI-governance rules"]),
        ("T", "Threats", "crosshair", "B71C1C",
         ["Slow, low-noise attackers", "Big vendors adding features", "Alert fatigue if mis-tuned"])]
sw4 = (12.98 - VX - 0.2 - 0.24) / 4
for i, (L, t, ic, col, items) in enumerate(swot):
    x = VX + 0.1 + i * (sw4 + 0.08)
    box(s, x, BY + 0.38, sw4, BH - 0.46, fill=CARD2, line=col, lw=1.5, radius=0.08)
    box(s, x, BY + 0.38, sw4, 0.36, fill=col, line=col, radius=0.08)
    text(s, x, BY + 0.38, sw4, 0.36, [[(t, {"bold": True})]], size=10.5, align=PP_ALIGN.CENTER,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, x + 0.04, BY + 0.78, sw4 - 0.06, BH - 0.9, items, size=9.5, after=3,
            bcolor=col if col != "B71C1C" else RED_HI)


# =================================================== 5. IMPACTS & BENEFITS
s = new_slide(5)
TY, TH = 1.72, 2.9
# top-left: impact on target audience (2x2)
box(s, 0.32, TY, 6.4, TH, fill=CARD, line=LINE, alpha=90)
text(s, 0.4, TY + 0.03, 6.24, 0.34, [[("Potential Impact on Target Audience", {"bold": True})]],
     size=13, font=HEAD, color=GOLD, align=PP_ALIGN.CENTER)
aud = [("shield-check", "Enterprise SOC Teams", "Fewer, clearer alerts that name the exact action and trigger."),
       ("bot", "AI / MLOps Teams", "Ship autonomous agents with a behavioural safety net."),
       ("landmark", "BFSI & Healthcare", "Stops silent data exfiltration from regulated systems."),
       ("building-2", "Govt & Defence IT", "Catches hijacked agents even with valid credentials.")]
qw, qh = 3.07, (TH - 0.5) / 2
for i, (ic, t, body) in enumerate(aud):
    x = 0.42 + (i % 2) * (qw + 0.06)
    y = TY + 0.42 + (i // 2) * (qh + 0.04)
    box(s, x, y, qw, qh - 0.04, fill=CARD2, line=RED_DK, radius=0.08)
    badge(s, ic, x + 0.1, y + (qh - 0.04 - 0.62) / 2, 0.62)
    text(s, x + 0.82, y + 0.08, qw - 0.88, 0.3, [[(t, {"bold": True})]], size=11, font=HEAD)
    text(s, x + 0.82, y + 0.38, qw - 0.88, qh - 0.5, body, size=10, color=MUTED)
arrow(s, 0.42 + qw + 0.03, TY + 0.48, 0.42 + qw + 0.03, TY + TH - 0.1, color=RED, w=1,
      head=False)

# top-right: impact on organisation
VX = 6.86
box(s, VX, TY, 12.98 - VX, TH, fill=CARD, line=LINE, alpha=90)
text(s, VX + 0.08, TY + 0.03, 12.98 - VX - 0.16, 0.34,
     [[("Impact on Target Organisation", {"bold": True})]], size=13, font=HEAD, color=GOLD,
     align=PP_ALIGN.CENTER)
cd_ = 1.7
ccx, ccy = VX + 0.15, TY + 0.42 + (TH - 0.5 - cd_) / 2
box(s, ccx, ccy, cd_, cd_, fill=RED, line="FFCDD2", lw=3, shape=MSO_SHAPE.OVAL)
icon(s, "radar", ccx + 0.55, ccy + 0.22, 0.6)
text(s, ccx, ccy + 0.84, cd_, 0.55, [[("SentinelTrace", {"bold": True})], "prototype"],
     size=12, font=HEAD, align=PP_ALIGN.CENTER, spacing=0.9)
org = [("Closes the blind spot", "flags hijacked agents whose tokens are still valid"),
       ("Stops injection exfiltration", "untrusted doc → sensitive action → external host"),
       ("Contains only the bad agent", "the rest of the fleet keeps working"),
       ("Audit-ready AI governance", "every score and signal is logged")]
oy = TY + 0.44
oh = (TH - 0.52) / 4
for i, (t, b) in enumerate(org):
    x = ccx + cd_ + 0.2 + (0.12 if i in (1, 2) else 0)
    w = 12.88 - x
    y = oy + i * oh
    arrow(s, ccx + cd_ - 0.04, ccy + cd_ / 2, x, y + (oh - 0.06) / 2, color=GOLD, w=1,
          head=False)
    box(s, x, y, w, oh - 0.06, fill="0D3B4A", line="26C6DA", lw=1.0, radius=(oh - 0.06) / 2)
    num_badge(s, i + 1, x + 0.04, y + (oh - 0.06 - 0.42) / 2, 0.42, fill="00838F", size=13)
    text(s, x + 0.52, y, w - 0.58, oh - 0.06,
         [[(t, {"bold": True})], [(b, {"size": 9, "color": "B2EBF2"})]], size=10.5,
         anchor=MSO_ANCHOR.MIDDLE, spacing=0.92)

# bottom-left: before vs after table
BY = TY + TH + 0.14
BH = 7.06 - BY
box(s, 0.32, BY, 6.4, BH, fill=CARD, line=LINE, alpha=90)
cols = [0.42, 2.12, 4.37]
cws = [1.66, 2.21, 2.25]
hdr = [("", None), ("Traditional IAM", "x"), ("SentinelTrace", "check")]
for j, (h, ic) in enumerate(hdr):
    if not h:
        continue
    fill = "4A1418" if j == 1 else "1B3A1E"
    ln = RED_HI if j == 1 else GREEN
    box(s, cols[j], BY + 0.08, cws[j], 0.36, fill=fill, line=ln, radius=0.08)
    icon(s, ic, cols[j] + 0.08, BY + 0.13, 0.26)
    text(s, cols[j] + 0.36, BY + 0.08, cws[j] - 0.4, 0.36, [[(h, {"bold": True})]], size=11,
         font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
rows = [("What it checks", "Token & permissions only", "Token + live behaviour"),
        ("Hijacked agent, valid token", "Passes silently", "Flagged in real time"),
        ("Alert quality", "None / black-box score", "Plain-language, per signal"),
        ("Response", "Revoke all, manually", "Log → approve → isolate"),
        ("Multi-agent chains", "Blind", "Risk propagates downstream")]
rh = (BH - 0.56) / len(rows)
for i, (a, b, c) in enumerate(rows):
    y = BY + 0.5 + i * rh
    if i % 2 == 0:
        box(s, 0.42, y, 6.2, rh, fill="2A0B0F", line=None, radius=0.04)
    text(s, cols[0], y, cws[0], rh, [[(a, {"bold": True})]], size=9.5, color=GOLD,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, cols[1] + 0.06, y, cws[1] - 0.06, rh, b, size=9.5, color="FFAB91",
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, cols[2] + 0.06, y, cws[2] - 0.06, rh, c, size=9.5, color="A5D6A7", bold=True,
         anchor=MSO_ANCHOR.MIDDLE)

# bottom-right: key benefits
box(s, VX, BY, 12.98 - VX, BH, fill=CARD, line=LINE, alpha=90)
text(s, VX, BY + 0.02, 12.98 - VX, 0.32, [[("Key Benefits", {"bold": True})]], size=13,
     font=HEAD, color=GOLD, align=PP_ALIGN.CENTER)
ben = [("bell-ring", "Cuts alert fatigue", "act on an alert, not triage it"),
       ("zap", "Real-time detection", "every action scored live"),
       ("workflow", "Keeps work moving", "only the bad agent is isolated"),
       ("file-check", "Audit-ready", "full trail for AI governance"),
       ("coins", "Low cost", "open-source, CPU-only"),
       ("puzzle", "Framework-agnostic", "one hook, any agent stack")]
kw = (12.98 - VX - 0.2 - 0.16) / 3
kh = (BH - 0.4 - 0.34 - 0.06) / 2
for i, (ic, t, b) in enumerate(ben):
    x = VX + 0.1 + (i % 3) * (kw + 0.08)
    y = BY + 0.38 + (i // 3) * (kh + 0.06)
    box(s, x, y, kw, kh, fill=CARD2, line=RED_DK, radius=0.08)
    badge(s, ic, x + 0.08, y + (kh - 0.44) / 2, 0.44)
    text(s, x + 0.58, y, kw - 0.62, kh, [[(t, {"bold": True})], [(b, {"size": 8.5, "color": MUTED})]],
         size=10, anchor=MSO_ANCHOR.MIDDLE, spacing=0.92)
text(s, VX + 0.1, BY + BH - 0.36, 12.98 - VX - 0.2, 0.3,
     [[("Aligned with  ", {"color": MUTED}),
       ("OWASP LLM Top 10  ·  NIST AI RMF  ·  DPDP Act 2023  ·  SDG 9 & 16", {"bold": True})]],
     size=9.5, color=GOLD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ================================================ 6. USP & FUTURE SCOPE
s = new_slide(6)
TY = 1.72
UW = 5.3
box(s, 0.32, TY, UW, 7.06 - TY, fill=CARD, line=LINE, alpha=92)
box(s, 0.32, TY, UW, 0.48, fill=RED_DK, line=LINE, radius=0.08)
badge(s, "gem", 0.42, TY + 0.06, 0.36, pad=0.2)
text(s, 0.86, TY, UW - 0.6, 0.48, [[("UNIQUE SELLING PROPOSITION", {"bold": True})]], size=14,
     font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
box(s, 0.45, TY + 0.6, UW - 0.26, 0.92, fill="2A0B0F", line=GOLD, lw=1, radius=0.08)
text(s, 0.55, TY + 0.6, UW - 0.46, 0.92,
     [[("Most tools ask ", {}), ("“is this token valid?”", {"bold": True, "color": "FFAB91"})],
      [("SentinelTrace asks ", {}),
       ("“is this agent still acting like its role?”", {"bold": True, "color": "A5D6A7"})]],
     size=10.5, italic=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, after=3)
usp = [("fingerprint", "Behaviour-aware identity", "live actions measured against a signed role contract"),
       ("share-2", "Chain-aware", "one compromised agent can't hide behind another's clean record"),
       ("message-square-warning", "Explainable by design", "every alert names the action, mismatch and trigger"),
       ("shield-check", "Graduated containment", "log → approve → isolate, never an all-or-nothing kill switch"),
       ("puzzle", "Open & pluggable", "100% open-source; swap any signal without a rebuild")]
uy = TY + 1.64
for ic, t, b in usp:
    badge(s, ic, 0.45, uy + 0.02, 0.44)
    text(s, 1.0, uy - 0.04, UW - 0.75, 0.28, [[(t, {"bold": True})]], size=11.5, font=HEAD,
         color=GOLD)
    text(s, 1.0, uy + 0.21, UW - 0.75, 0.3, b, size=9.5, color=CREAM)
    uy += 0.56
# chain-aware mini diagram
dy = uy + 0.06
box(s, 0.45, dy, UW - 0.26, 7.0 - dy, fill="1E0B0D", line=RED_DK, radius=0.08)
for i, (lab, sc, col) in enumerate([("A", 88, RED_HI), ("B", 53, AMBER), ("C", 32, GOLD)]):
    x = 0.75 + i * 1.65
    c = box(s, x, dy + 0.1, 0.46, 0.46, fill=col, line=CREAM, lw=1.25, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = lab
    r.font.size, r.font.bold, r.font.name = Pt(14), True, HEAD
    r.font.color.rgb = rgb("1A0507")
    text(s, x + 0.5, dy + 0.12, 0.8, 0.42, [[(f"risk {sc}", {"bold": True})]], size=10,
         color=col, anchor=MSO_ANCHOR.MIDDLE)
    if i < 2:
        arrow(s, x + 1.15, dy + 0.33, x + 1.62, dy + 0.33, color=GOLD, w=2)
text(s, 0.5, dy + 0.58, UW - 0.36, 0.26, "A calls B calls C — each inherits a share of upstream risk",
     size=9, italic=True, color=MUTED, align=PP_ALIGN.CENTER)

# right box: future scope roadmap
FX = 0.32 + UW + 0.14
FW = 12.98 - FX
box(s, FX, TY, FW, 7.06 - TY, fill=CARD, line=LINE, alpha=92)
box(s, FX, TY, FW, 0.48, fill=RED_DK, line=LINE, radius=0.08)
badge(s, "route", FX + 0.1, TY + 0.06, 0.36, pad=0.2)
text(s, FX + 0.54, TY, FW - 0.6, 0.48,
     [[("FUTURE SCOPE", {"bold": True}), ("  —  from hackathon to platform", {"size": 11, "italic": True, "color": MUTED})]],
     size=14, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
phases = [("flag", "Hackathon", "24 hrs", "3–5 simulated agents · 3 fused signals + rule layer · live Streamlit dashboard",
           GREEN),
          ("search-check", "Pilot", "Month 1–3", "1 real agent fleet · tune thresholds on real data · target < 5% false-positive rate",
           "26A69A"),
          ("webhook", "Integrate", "Month 4–6", "MCP & A2A tool-call hooks · LangChain / AutoGen adapters · ready-made OPA policy packs",
           AMBER),
          ("badge-check", "Standardise", "Month 7–12", "Adopt emerging agent-identity standards (OAuth, SPIFFE-style workload IDs) instead of a custom role format",
           "FB8C00"),
          ("globe", "Federate", "Year 2", "Cross-org threat intel: share drift patterns, never raw logs · SaaS + on-prem editions",
           RED_HI)]
py = TY + 0.62
ph = 0.72
line_x = FX + 1.62
arrow(s, line_x, py + 0.2, line_x, py + 4 * ph + 0.3, color=GOLD, w=2.5)
for i, (ic, t, when, body, col) in enumerate(phases):
    y = py + i * ph
    pill(s, FX + 0.12, y + 0.14, 1.22, 0.3, when, fill=CARD2, line=col, color=col, size=9.5)
    box(s, line_x - 0.24, y + 0.05, 0.48, 0.48, fill=col, line=CREAM, lw=1.5, shape=MSO_SHAPE.OVAL)
    icon(s, ic, line_x - 0.14, y + 0.15, 0.28)
    box(s, line_x + 0.36, y, FW - (line_x - FX) - 0.46, ph - 0.08, fill=CARD2, line=col, lw=1,
        radius=0.08)
    text(s, line_x + 0.46, y + 0.01, FW - (line_x - FX) - 0.6, ph - 0.1,
         [[(f"Phase {i}  ·  ", {"color": col, "bold": True}), (t, {"bold": True, "font": HEAD})],
          [(body, {"size": 9.5, "color": MUTED})]],
         size=11, anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
my = py + 5 * ph + 0.06
mh = 7.0 - my
targets = [("< 1 s", "scoring per action"), ("< 5 %", "false positives post-pilot"),
           ("100 %", "alerts explained"), ("0", "raw logs shared")]
mw = (FW - 0.2 - 0.18) / 4
text(s, FX + 0.1, my - 0.02, 2, 0.22, [[("TARGET METRICS", {"bold": True})]], size=8.5,
     color=GOLD)
for i, (v, l) in enumerate(targets):
    x = FX + 0.1 + i * (mw + 0.06)
    box(s, x, my + 0.2, mw, mh - 0.2, fill="2A0B0F", line=GOLD, lw=1, radius=0.08)
    text(s, x, my + 0.2, mw, mh - 0.2, [[(v, {"bold": True, "size": 16, "font": HEAD, "color": GOLD})],
                                        [(l, {"size": 9, "color": MUTED})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=0.9)

prs.core_properties.title = "D3COY 1.0 – CTRL FREAKS – Identity Drift Detection in Autonomous AI Agents"
prs.core_properties.author = "CTRL FREAKS"
prs.save(OUT)
print("saved", OUT)
